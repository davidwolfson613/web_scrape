import requests
import sys
import urllib3
from bs4 import BeautifulSoup
from tabulate import tabulate
from docx import Document
import time

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

def get_hb():
  
  '''
  This function obtains the HB number (identification number) of equipment of which calibration data is needed. The user inputs the HB number(s).
  '''
  
  hb = []
  start = input('Start Date: ')
  end = input('End Date: ')
  loc = input('Location: ')
  tmp = input('Enter HB number of equipment (type "done" when finished): ')

  while tmp.lower() != 'done':
    hb.append(tmp.upper())
    tmp = input('Enter HB number of equipment (type "done" when finished): ')

  return hb,start,end,loc

def get_cal_dates(hb_lst,start,end,loc):

  '''
  This function automatically obtains all the calibration data from the internet in the form of HTML data. The HTML data is then parsed and the relevant data is stored.
  '''
  cal_dates = []
  manufact = []
  desc = []
  start_lst = [start]*len(hb_lst)
  end_lst = [end]*len(hb_lst)
  loc_lst = [loc]*len(hb_lst)

  for hb in hb_lst:
    
    # the url containing the calibration data
    url = 'https://geets.gm.com/webapps/metex1/met_ex.exe?INPUTFORM=RUNREPORT&REPORTNAME=INVENTORY%3A+BY+EQUIPMENT+INFORMATION&PROMPT1=%25&PROMPT2=%25&PROMPT3=%25&PROMPT4=%25&PROMPT5=%25&PROMPT6=%25&PROMPT7=%25&PROMPT8=%25&PROMPT9=%25&PROMPT10={}&PROMPT11=%25&PROMPT12=%25'.format(hb)
    
    # use requests to get the HTML data
    html = requests.get(url,verify=False)
    time.sleep(3)
  
    # use BeautifulSoup to parse HTML for relevant data
    soup = BeautifulSoup(html.text, 'html.parser')
    table = soup.find(id='sqlreport')

    try:
      rows = table.find_all('tr')
    except AttributeError:
      print(f'Trouble processing {hb}. Most likely a timeout. Moving on to next HB number.')
      continue

    print(f'Processing {hb}')
    for row in rows:
        
        columns_h = row.find_all('th')
        columns_d = row.find_all('td')

        if len(columns_h) > 0:

            str_cols_h = [str(i) for i in columns_h]
            idx = str_cols_h.index('<th>Due-Date</th>')
            idx1 = str_cols_h.index('<th>Manufacturer</th>')
            idx2 = str_cols_h.index('<th>Description</th>')

        if len(columns_d) > idx:

            tmp = str(columns_d[idx]).split('<td>')[1].split('</td>')[0]
            cal_date = tmp if tmp != '\xa0' else 'N/A'
            cal_dates.append(cal_date)
            manufact.append(str(columns_d[idx1]).split('<td>')[1].split('</td>')[0])
            desc.append(str(columns_d[idx2]).split('<td>')[1].split('</td>')[0])
  
  # create dictionary containing calibration data
  data_dict = {'Equipment Number':hb_lst,'Location':loc_lst,'Manufacturer':manufact,'Description':desc,
                  'Cal Date':cal_dates,'Start Date':start_lst,'End Date':end_lst}
  return data_dict

def make_table(data_dict):

  '''
  This function creates a table in Microsoft Word from the calibration data dictionary
  '''
  
  # create the Word doc
  doc = Document()
  
  # add a table for the calibration data
  table = doc.add_table(rows=len(data_dict['Equipment Number'])+1,cols=len(data_dict))
  
  # fill table with data
  for i,name in enumerate(data_dict.keys()):
    for j in range(len(table.rows)):
      if j==0:
        table.cell(j,i).text = name
      else:
        table.cell(j,i).text = data_dict[name][j-1] if data_dict[name][j-1] != '' else 'N/A'
  
  # style the table and save it
  table.style = 'Table Grid'
  doc.save('Equipment table.docx')

if __name__ == '__main__':

  hb,start,end,loc = get_hb()
  data = get_cal_dates(hb,start,end,loc)
  make_table(data)
